"""
HSBLearn Document Processing & Text Chunking Service.
Handles PDF, DOCX, PPTX, and TXT files.
Extracts page and slide metadata for RAG grounding.
"""
import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
from backend.database import get_db

# Chunking configuration
CHUNK_SIZE_WORDS = 250
CHUNK_OVERLAP_WORDS = 40

def chunk_text(text: str, page_or_slide: Optional[int] = None) -> List[Dict[str, Any]]:
    """
    Splits text into overlapping chunks while preserving page/slide numbering.
    """
    cleaned = re.sub(r'\s+', ' ', text).strip()
    if not cleaned:
        return []

    words = cleaned.split(' ')
    chunks = []
    chunk_index = 0

    start = 0
    while start < len(words):
        end = min(start + CHUNK_SIZE_WORDS, len(words))
        chunk_str = " ".join(words[start:end])
        if len(chunk_str.strip()) > 10:
            chunks.append({
                "chunk_index": chunk_index,
                "page_or_slide": page_or_slide,
                "text": chunk_str
            })
            chunk_index += 1
        if end >= len(words):
            break
        start += (CHUNK_SIZE_WORDS - CHUNK_OVERLAP_WORDS)

    return chunks

def extract_pdf(file_path: str) -> List[Dict[str, Any]]:
    import pypdf
    chunks = []
    try:
        reader = pypdf.PdfReader(file_path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("PDF is password encrypted.")

        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            page_chunks = chunk_text(text, page_or_slide=page_num)
            chunks.extend(page_chunks)
    except Exception as e:
        raise ValueError(f"Failed to extract PDF: {str(e)}")
    return chunks

def extract_docx(file_path: str) -> List[Dict[str, Any]]:
    import docx
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        return chunk_text(full_text, page_or_slide=None)
    except Exception as e:
        raise ValueError(f"Failed to extract DOCX: {str(e)}")

def extract_pptx(file_path: str) -> List[Dict[str, Any]]:
    from pptx import Presentation
    chunks = []
    try:
        prs = Presentation(file_path)
        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
            slide_content = " ".join(slide_texts)
            chunks.extend(chunk_text(slide_content, page_or_slide=slide_num))
    except Exception as e:
        raise ValueError(f"Failed to extract PPTX: {str(e)}")
    return chunks

def extract_txt(file_path: str) -> List[Dict[str, Any]]:
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return chunk_text(text, page_or_slide=None)
    except Exception as e:
        raise ValueError(f"Failed to extract TXT: {str(e)}")

def process_material_file(material_id: int):
    """
    Extracts text chunks from the material file and saves them into extracted_chunks table.
    Updates materials table with extraction_status.
    """
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, file_name, file_type, storage_path FROM materials WHERE id = ?", (material_id,))
        mat = cursor.fetchone()
        if not mat:
            return

        cursor.execute("UPDATE materials SET extraction_status = 'processing' WHERE id = ?", (material_id,))

    file_path = mat["storage_path"]
    ext = mat["file_type"].lower()
    if not ext.startswith("."):
        ext = f".{ext}"

    chunks: List[Dict[str, Any]] = []
    status = "ready"
    status_detail = "Extracted successfully"

    if not os.path.exists(file_path):
        status = "failed"
        status_detail = "Physical file missing on server."
    else:
        try:
            if ext == ".pdf":
                chunks = extract_pdf(file_path)
            elif ext in [".docx", ".doc"]:
                chunks = extract_docx(file_path)
            elif ext in [".pptx", ".ppt"]:
                chunks = extract_pptx(file_path)
            elif ext in [".txt", ".md"]:
                chunks = extract_txt(file_path)
            else:
                status = "unavailable"
                status_detail = f"File format {ext} extraction unsupported."

            if status == "ready" and len(chunks) == 0:
                status = "unavailable"
                status_detail = "No extractable text found (scanned image or empty document)."
        except Exception as e:
            status = "failed"
            status_detail = str(e)

    with get_db() as conn:
        cursor = conn.cursor()
        # Delete prior chunks if re-processing
        cursor.execute("DELETE FROM extracted_chunks WHERE material_id = ?", (material_id,))

        if status == "ready" and chunks:
            for idx, c in enumerate(chunks):
                cursor.execute("""
                    INSERT INTO extracted_chunks (material_id, page_or_slide, chunk_index, text)
                    VALUES (?, ?, ?, ?)
                """, (material_id, c["page_or_slide"], idx, c["text"]))

        cursor.execute("""
            UPDATE materials 
            SET extraction_status = ?, extracted_text_status = ? 
            WHERE id = ?
        """, (status, status_detail, material_id))
